"""
Утилиты для экспорта отчетов
"""

import os
import logging
from datetime import datetime
from typing import List, Dict, Any, Optional

from PyQt6.QtWidgets import QFileDialog, QMessageBox


class ExportService:
    """Сервис для экспорта отчетов в различные форматы"""
    
    def __init__(self):
        self.logger = logging.getLogger(__name__)
    
    def export_to_excel(self, data: List[Dict[str, Any]], filename: str, sheet_name: str = "Report") -> bool:
        """Экспорт данных в Excel"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
            
            wb = Workbook()
            ws = wb.active
            ws.title = sheet_name
            
            if not data:
                return False
            
            headers = list(data[0].keys())
            
            header_fill = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header.replace('_', ' ').title())
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = thin_border
            
            for row_idx, row_data in enumerate(data, 2):
                for col_idx, key in enumerate(headers, 1):
                    value = row_data.get(key, '')
                    if isinstance(value, (int, float)):
                        cell = ws.cell(row=row_idx, column=col_idx, value=value)
                        cell.alignment = Alignment(horizontal='right')
                    else:
                        cell = ws.cell(row=row_idx, column=col_idx, value=str(value) if value else '')
                        cell.alignment = Alignment(horizontal='left')
                    cell.border = thin_border
            
            for col in range(1, len(headers) + 1):
                ws.column_dimensions[chr(64 + col) if col <= 26 else 'A' + chr(64 + col - 26)].width = 20
            
            wb.save(filename)
            self.logger.info(f"Экспорт в Excel: {filename}")
            return True
            
        except ImportError:
            self.logger.error("openpyxl не установлен")
            return False
        except Exception as e:
            self.logger.error(f"Ошибка экспорта в Excel: {e}")
            return False
    
    def export_to_pdf(self, title: str, data: List[Dict[str, Any]], filename: str) -> bool:
        """Экспорт данных в PDF"""
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
            
            doc = SimpleDocTemplate(filename, pagesize=landscape(A4))
            elements = []
            styles = getSampleStyleSheet()
            
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=18,
                textColor=colors.HexColor('#4CAF50'),
                spaceAfter=20
            )
            
            elements.append(Paragraph(title, title_style))
            elements.append(Spacer(1, 0.2 * inch))
            
            if data:
                headers = list(data[0].keys())
                table_data = [headers]
                
                for row in data:
                    table_data.append([str(row.get(h, '')) for h in headers])
                
                table = Table(table_data)
                table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4CAF50')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 10),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#f5f5f5')),
                    ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#cccccc')),
                    ('FONTSIZE', (0, 1), (-1, -1), 8),
                ]))
                
                elements.append(table)
            
            doc.build(elements)
            self.logger.info(f"Экспорт в PDF: {filename}")
            return True
            
        except ImportError:
            self.logger.error("reportlab не установлен")
            return False
        except Exception as e:
            self.logger.error(f"Ошибка экспорта в PDF: {e}")
            return False
    
    def export_to_word(self, title: str, data: List[Dict[str, Any]], filename: str) -> bool:
        """Экспорт данных в Word"""
        try:
            from docx import Document
            from docx.shared import Inches, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.enum.table import WD_TABLE_ALIGNMENT
            
            doc = Document()
            
            title_heading = doc.add_heading(title, 0)
            title_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in title_heading.runs:
                run.font.color.rgb = RGBColor(76, 175, 80)
            
            doc.add_paragraph()
            
            if data:
                headers = list(data[0].keys())
                table = doc.add_table(rows=1, cols=len(headers))
                table.style = 'Light Grid Accent 1'
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                header_cells = table.rows[0].cells
                for i, header in enumerate(headers):
                    header_cells[i].text = header.replace('_', ' ').title()
                    for paragraph in header_cells[i].paragraphs:
                        for run in paragraph.runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                
                for row_data in data:
                    row_cells = table.add_row().cells
                    for i, key in enumerate(headers):
                        row_cells[i].text = str(row_data.get(key, ''))
            
            doc.add_paragraph()
            doc.add_paragraph(f"Дата создания: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
            
            doc.save(filename)
            self.logger.info(f"Экспорт в Word: {filename}")
            return True
            
        except ImportError:
            self.logger.error("python-docx не установлен")
            return False
        except Exception as e:
            self.logger.error(f"Ошибка экспорта в Word: {e}")
            return False
    
    def export_to_csv(self, data: List[Dict[str, Any]], filename: str) -> bool:
        """Экспорт данных в CSV"""
        try:
            import csv
            
            if not data:
                return False
            
            headers = list(data[0].keys())
            
            with open(filename, 'w', newline='', encoding='utf-8-sig') as f:
                writer = csv.DictWriter(f, fieldnames=headers)
                writer.writeheader()
                writer.writerows(data)
            
            self.logger.info(f"Экспорт в CSV: {filename}")
            return True
            
        except Exception as e:
            self.logger.error(f"Ошибка экспорта в CSV: {e}")
            return False
    
    def export_to_html(self, title: str, data: List[Dict[str, Any]], filename: str) -> bool:
        """Экспорт данных в HTML"""
        try:
            if not data:
                return False
            
            headers = list(data[0].keys())
            
            html = f"""
<!DOCTYPE html>
<html>
<head>
    <meta charset="UTF-8">
    <title>{title}</title>
    <style>
        body {{ font-family: Arial, sans-serif; margin: 20px; }}
        h1 {{ color: #4CAF50; text-align: center; }}
        table {{ border-collapse: collapse; width: 100%; margin-top: 20px; }}
        th {{ background-color: #4CAF50; color: white; padding: 12px; text-align: left; }}
        td {{ border: 1px solid #ddd; padding: 8px; }}
        tr:nth-child(even) {{ background-color: #f5f5f5; }}
        tr:hover {{ background-color: #e8f5e9; }}
    </style>
</head>
<body>
    <h1>{title}</h1>
    <p>Дата создания: {datetime.now().strftime('%d.%m.%Y %H:%M')}</p>
    <table>
        <tr>
"""
            
            for header in headers:
                html += f"            <th>{header.replace('_', ' ').title()}</th>\n"
            
            html += """        </tr>
"""
            
            for row in data:
                html += "        <tr>\n"
                for key in headers:
                    value = row.get(key, '')
                    html += f"            <td>{value}</td>\n"
                html += "        </tr>\n"
            
            html += """    </table>
</body>
</html>
"""
            
            with open(filename, 'w', encoding='utf-8') as f:
                f.write(html)
            
            self.logger.info(f"Экспорт в HTML: {filename}")
            return True
            
        except Exception as e:
            self.logger.error(f"Ошибка экспорта в HTML: {e}")
            return False
    
    def get_save_filename(self, parent, title: str, format_type: str) -> Optional[str]:
        """Получить имя файла для сохранения"""
        filters = {
            'excel': "Excel Files (*.xlsx);;All Files (*)",
            'pdf': "PDF Files (*.pdf);;All Files (*)",
            'word': "Word Files (*.docx);;All Files (*)",
            'csv': "CSV Files (*.csv);;All Files (*)",
            'html': "HTML Files (*.html);;All Files (*)"
        }
        
        default_ext = {
            'excel': '.xlsx',
            'pdf': '.pdf',
            'word': '.docx',
            'csv': '.csv',
            'html': '.html'
        }
        
        filename, _ = QFileDialog.getSaveFileName(
            parent,
            title,
            f"report_{datetime.now().strftime('%Y%m%d_%H%M')}{default_ext.get(format_type, '')}",
            filters.get(format_type, "All Files (*)")
        )
        
        return filename


def show_export_dialog(parent, api_service, report_type: str, data_func):
    """Показать диалог экспорта"""
    export_service = ExportService()
    
    format_options = {
        'excel': ('Экспорт в Excel', 'excel'),
        'pdf': ('Экспорт в PDF', 'pdf'),
        'word': ('Экспорт в Word', 'word'),
        'csv': ('Экспорт в CSV', 'csv'),
        'html': ('Экспорт в HTML', 'html')
    }
    
    buttons_html = """
    <div style="text-align: center; margin: 10px;">
        <p>Выберите формат:</p>
"""
    
    for key, (text, fmt) in format_options.items():
        buttons_html += f'<button onclick="export_{key}()" style="padding: 10px 20px; margin: 5px; cursor: pointer; background: #4CAF50; color: white; border: none; border-radius: 4px;">{text}</button>'
    
    buttons_html += """
    </div>
    """
    
    reply = QMessageBox(parent)
    reply.setWindowTitle("Экспорт отчета")
    reply.setText(f"Выберите формат для экспорта отчета '{report_type}':")
    reply.setIcon(QMessageBox.Icon.Information)
    
    msg = QMessageBox.question(
        parent,
        "Экспорт отчета",
        "Выберите формат для экспорта:",
        QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
    )
    
    if msg == QMessageBox.StandardButton.Yes:
        from PyQt6.QtWidgets import QInputDialog
        formats = ['Excel (.xlsx)', 'PDF (.pdf)', 'Word (.docx)', 'CSV (.csv)', 'HTML (.html)']
        format_map = {'Excel (.xlsx)': 'excel', 'PDF (.pdf)': 'pdf', 'Word (.docx)': 'word', 'CSV (.csv)': 'csv', 'HTML (.html)': 'html'}
        
        selected, ok = QInputDialog.getItem(parent, "Формат", "Выберите формат:", formats, 0, False)
        
        if ok and selected:
            fmt = format_map[selected]
            filename = export_service.get_save_filename(parent, "Сохранить отчет", fmt)
            
            if filename:
                data = data_func()
                success = False
                
                if fmt == 'excel':
                    success = export_service.export_to_excel(data, filename)
                elif fmt == 'pdf':
                    success = export_service.export_to_pdf("Отчет", data, filename)
                elif fmt == 'word':
                    success = export_service.export_to_word("Отчет", data, filename)
                elif fmt == 'csv':
                    success = export_service.export_to_csv(data, filename)
                elif fmt == 'html':
                    success = export_service.export_to_html("Отчет", data, filename)
                
                if success:
                    QMessageBox.information(parent, "Успех", f"Отчет сохранен:\n{filename}")
                else:
                    QMessageBox.warning(parent, "Ошибка", "Не удалось сохранить отчет")